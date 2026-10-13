# -*- coding: utf-8 -*-
"""CSV / TXT / Excel exports: the right rows, at the right rate.

2026-09-02, reported against the CSV connection dialog: "when setting a
parallel database logging we should be able to pick different gateways
running, selecting them, and all the tags we want to export, and also the
rate. At the moment this is not working."

It was not working for five separate reasons, and each one is asserted here.

  1. The gateway chips and tag list were built from gateway_configurations
     ALONE. The reporting install has none - its only source is a power meter
     under power_management_config, producing forty tags a second - so both
     lists were empty and there was nothing to pick. Asserted on the frontend
     source, since that is where the list is built.
  2. A meter only ever wrote to the connection named in its OWN database_id.
     Ticking the meter in a CSV sink's gateway list did nothing, because
     nothing ever read that list on the meter path.
  3. The meter path ignored tag_filters and gateway_filters outright, so a
     CSV scoped to two tags still received all forty registers.
  4. There was no export RATE anywhere. A 1 s meter with 40 registers writes
     40 rows a second - 3.5 million a day - into a file meant to be read.
  5. There was no Excel engine at all.

And the trap that would have silently undone all of it: the connection is
saved through a hand-written field list, and sent to the worker through a
SECOND one. A field missing from either is dropped without a word, which
looks precisely like the feature not working.
"""
from __future__ import annotations

import io
import os
import sys
import tempfile
from datetime import datetime

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "backend"))
os.environ.setdefault("TRUSTNODE_SKIP_DOTENV", "1")

FAILS: list[str] = []
NL = chr(10)          # written literally so no escape survives an edit


def check(name, ok, detail=""):
    print("  {0:58s}: {1}{2}".format(name, "PASS" if ok else "FAIL",
                                     (" - " + str(detail)[:150]) if detail else ""))
    if not ok:
        FAILS.append(name)


from app.services import file_sinks as fs          # noqa: E402
from app.services.power_manager import PowerManager  # noqa: E402

print("TrustNode - CSV / TXT / Excel exports")
print()
print("[the sink exports what was ticked, and nothing else]")
ROWS = [
    {"gateway_id": "EM1", "tag_name": "current_l1_a", "value": 0.24, "ts_utc": "t1"},
    {"gateway_id": "EM1", "tag_name": "voltage_l1_v", "value": 239.6, "ts_utc": "t1"},
    {"gateway_id": "gw-2", "tag_name": "current_l1_a", "value": 7.0, "ts_utc": "t1"},
]
check("no filters means everything, as before",
      len(fs.filter_rows({}, ROWS)) == 3,
      "a config written before scoping existed must keep working")
check("a gateway filter keeps only that source",
      [r["gateway_id"] for r in fs.filter_rows({"gateway_filters": ["EM1"]}, ROWS)]
      == ["EM1", "EM1"])
check("a tag filter keeps only those tags",
      [r["tag_name"] for r in fs.filter_rows({"tag_filters": ["current_l1_a"]}, ROWS)]
      == ["current_l1_a", "current_l1_a"])
check("  both together are an AND",
      len(fs.filter_rows({"gateway_filters": ["EM1"],
                          "tag_filters": ["current_l1_a"]}, ROWS)) == 1,
      "the meter used to write all forty registers whatever was ticked")

print()
print("[the export rate]")
fs.reset_rate_state()
check("no rate configured writes every row",
      len(fs.due_rows("s0", {}, ROWS)) == 3,
      "an existing sink must not suddenly start throttling")

fs.reset_rate_state()
sink = {"export_interval_s": 60}
first = fs.due_rows("s1", sink, ROWS, now=1000.0)
held = fs.due_rows("s1", sink, ROWS, now=1030.0)
later = fs.due_rows("s1", sink, ROWS, now=1061.0)
check("the first cycle writes immediately", len(first) == 3)
check("  inside the interval nothing is written", held == [],
      "%d row(s) leaked" % len(held))
check("  and the next interval writes again", len(later) == 3)

fs.reset_rate_state()
fs.due_rows("s2", sink, [{"gateway_id": "EM1", "tag_name": "a", "value": 1}], now=0.0)
fs.due_rows("s2", sink, [{"gateway_id": "EM1", "tag_name": "a", "value": 2}], now=10.0)
out = fs.due_rows("s2", sink, [{"gateway_id": "EM1", "tag_name": "a", "value": 3}], now=61.0)
check("a throttled tag reports its LATEST value",
      len(out) == 1 and out[0]["value"] == 3,
      "got %s - an export sampled at 60 s must show what the tag reads now, "
      "not what it read 59 s ago" % [r.get("value") for r in out])

fs.reset_rate_state()
many = [{"gateway_id": "EM1", "tag_name": "t%d" % i, "value": i} for i in range(40)]
fs.due_rows("s3", sink, many, now=0.0)
for step in range(1, 60):
    fs.due_rows("s3", sink, many, now=float(step))
released = fs.due_rows("s3", sink, many, now=61.0)
check("one row per tag per interval, not one per cycle",
      len(released) == 40,
      "%d rows - 40 tags at 1 s unthrottled is 3.5 million rows a day" % len(released))

print()
print("[an Excel workbook rolls over]")
check("daily rollover names the file by day",
      fs.rollover_path("/x/power.xlsx", "daily", datetime(2026, 9, 2, 14, 5))
      == os.path.join("/x", "power_2026-09-02.xlsx").replace("\\", "/")
      or fs.rollover_path("/x/power.xlsx", "daily", datetime(2026, 9, 2, 14, 5)).endswith("power_2026-09-02.xlsx"),
      fs.rollover_path("/x/power.xlsx", "daily", datetime(2026, 9, 2, 14, 5)))
check("  hourly splits by hour",
      fs.rollover_path("/x/p.xlsx", "hourly", datetime(2026, 9, 2, 14, 5)).endswith("p_2026-09-02_14.xlsx"))
check("  monthly splits by month",
      fs.rollover_path("/x/p.xlsx", "monthly", datetime(2026, 9, 2)).endswith("p_2026-09.xlsx"))
check("  'none' keeps the configured path",
      fs.rollover_path("/x/p.xlsx", "none") == "/x/p.xlsx",
      "an operator who wants one file must be able to have one file")

tmp = tempfile.mkdtemp(prefix="tn-xlsx-")
target = os.path.join(tmp, "export.xlsx")
rows_a = [{"gateway_id": "EM1", "tag_name": "current_l1_a", "value": 0.24,
           "ts_utc": "2026-09-02 10:00:00", "quality": 192, "quality_label": "GOOD",
           "source": "power_modbus", "gateway_name": "EM1"}]
written = fs.write_xlsx(target, rows_a, "daily", lambda t: t,
                        now=datetime(2026, 9, 2, 10, 0))
check("the workbook is created", bool(written) and os.path.exists(written), written)

from openpyxl import load_workbook                  # noqa: E402
book = load_workbook(written)
sheet = book.active
check("  with a header row", [c.value for c in sheet[1]] == fs.CANONICAL_HEADER)
check("  and the reading", sheet.cell(row=2, column=5).value == "current_l1_a")

fs.write_xlsx(target, rows_a, "daily", lambda t: t, now=datetime(2026, 9, 2, 11, 0))
sheet2 = load_workbook(written).active
check("  appending adds a row rather than replacing the file",
      sheet2.max_row == 3, "max_row=%d" % sheet2.max_row)

other = fs.write_xlsx(target, rows_a, "daily", lambda t: t,
                      now=datetime(2026, 9, 3, 10, 0))
check("the next day is a NEW workbook", other != written and os.path.exists(other),
      "one ever-growing .xlsx is rewritten in full on every save")
check("  no half-written .part file is left behind",
      not any(f.endswith(".part") for f in os.listdir(tmp)),
      "the save writes to a temp file and renames, so a crash cannot destroy "
      "the previous good workbook")

print()
print("[a meter is reachable as an export source]")
src = io.open(os.path.join(ROOT, "backend", "app", "services", "power_manager.py"),
              encoding="utf-8", errors="replace").read()
check("a sink that NAMES the meter is honoured",
      "_sinks_for_device" in src and "gateway_filters" in src,
      "the meter only ever wrote to its own database_id, so ticking it in a "
      "CSV connection did nothing at all")
check("  the meter path applies the filters",
      "_file_sinks.filter_rows" in src)
check("  and the export rate", "_file_sinks.due_rows" in src)
check("  and can write a workbook", "write_xlsx" in src)

plc = io.open(os.path.join(ROOT, "backend", "app", "services", "plc_manager.py"),
              encoding="utf-8", errors="replace").read()
check("the PLC path throttles too", "_file_sinks.due_rows" in plc,
      "one rate everywhere - two implementations would drift")
check("  and writes xlsx", "_persist_xlsx_file_for_sink" in plc)

print()
print("[the meter is HANDED its sinks - it does not read a legacy table]")
# 2026-09-03, the reason the CSV export never wrote a single row. The meter
# resolved its own connections via app_store.get_bootstrap(), which reads
# `config_documents` - the LEGACY unscoped table. Connections are saved to
# config_documents_scoped. On the reporting install the legacy copy was last
# written 2026-06-19 and held only "Local SQLite", so the CSV export was
# invisible and BOTH routes to it were dead: gateway_filters and database_id.
#
# A PLC gateway never had this problem because the UI hands it resolved sinks.
# The meter now works the same way.
pm_obj = PowerManager.__new__(PowerManager)
SINK = {"id": "db-1", "name": "CSV Mirror", "engine": "csv_file",
        "file_path": "x.csv", "gateway_filters": ["EM1"], "export_interval_s": 1}
got = pm_obj._sinks_for_device("EM1", {"id": "EM1", "sinks": [SINK]})
check("the pushed sinks are used as given",
      [s.get("id") for s in got] == ["db-1"],
      "%s - the UI resolved these; the meter must not look them up again" % got)
check("  and NO config read happens for them",
      pm_obj._sinks_for_device("EM1", {"id": "EM1", "sinks": []}) == [],
      "an empty pushed list means 'no sinks', not 'go and guess'")
check("  a malformed entry is ignored, not written to",
      pm_obj._sinks_for_device("EM1", {"id": "EM1", "sinks": [{"no": "engine"}]}) == [])

src_pm = io.open(os.path.join(ROOT, "backend", "app", "services", "power_manager.py"),
                 encoding="utf-8", errors="replace").read()
body = src_pm.split("def _sinks_for_device")[1].split("def _utc_str_to_local_iso")[0]
# Match an actual CALL, not the word: the method's own docstring explains the
# bug and names get_bootstrap, and a substring check on prose is not a test.
code_lines = [ln for ln in body.splitlines() if not ln.strip().startswith("#")]
check("_sinks_for_device no longer calls get_bootstrap",
      not any("self._app_store.get_bootstrap(" in ln for ln in code_lines),
      "get_bootstrap reads config_documents, which the connection saver does "
      "not write - that is the whole bug")

print()
print("[the UI can offer a meter as a source]")
app = io.open(os.path.join(ROOT, "frontend", "src", "App.jsx"),
              encoding="utf-8", errors="replace").read()
# NOT named exportSources: api.js exports a FUNCTION by that name, imported
# into App.jsx and handed to the Data Export page. A const with the same name
# shadowed it, so that page called an array and every one of its filters came
# up empty - reported on 2026-09-29 as "the filters are not working".
check("the export scope is built from gateways AND meters",
      "const exportScopeSources" in app and "powerConfig?.devices" in app,
      "it read gateway_configurations alone, and the reporting install had "
      "none - so both pickers were empty while a meter produced 40 tags/s")
check("  the scope dialog uses that list, not gatewayConfigs",
      "for (const g of exportScopeSources)" in app
      and "{exportScopeSources.map((g)" in app)
check("  and it does not shadow api.js's exportSources()",
      "const exportSources" not in app,
      "the Data Export page receives `exportSources` as a callable")
check("Excel is offered as an engine", 'value="xlsx_file"' in app)
check("a file export can be created from the device picker",
      '__new_csv__' in app and '__new_xlsx__' in app,
      "requested: selecting a new one opens a popup to configure the file")

print()
print("[wide layout: one row per timestamp, a column per tag]")
# 2026-09-03, requested: "the first column is the timestamp and the columns
# different tags, so we can get multiple values per timestamp".
WROWS = [
    {"gateway_id": "EM1", "ts_utc": "t1", "tag_name": "voltage_v", "value": 230.0},
    {"gateway_id": "EM1", "ts_utc": "t1", "tag_name": "current_a", "value": 1.5},
    {"gateway_id": "EM1", "ts_utc": "t2", "tag_name": "voltage_v", "value": 231.0},
]
wide = {"layout": "wide"}
check("the header starts with the timestamp",
      fs.header_for(wide, WROWS)[0] == "ts_local", fs.header_for(wide, WROWS))
check("  and carries one column per tag",
      fs.header_for(wide, WROWS)[3:] == ["current_a", "voltage_v"],
      fs.header_for(wide, WROWS))
shaped = fs.shape_rows(wide, WROWS, lambda t: t)
check("both tags of one timestamp share a row",
      shaped[0] == ["t1", "t1", "EM1", 1.5, 230.0], shaped[0])
check("  a tag missing that cycle is BLANK, not zero",
      shaped[1][3] == "" and shaped[1][4] == 231.0,
      "%s - a blank cell and a real zero must not look the same" % shaped[1])
check("tag_filters fixes the column order",
      fs.header_for({"layout": "wide", "tag_filters": ["voltage_v", "current_a"]},
                    WROWS)[3:] == ["voltage_v", "current_a"],
      "the operator's own order, stable across restarts")
check("long stays the default", fs.layout_of({}) == "long",
      "every existing sink must keep the layout it has")

print()
print("[a file whose header does not match is rotated, not appended to]")
rt = tempfile.mkdtemp(prefix="tn-rot-")
target = os.path.join(rt, "export.csv")
with io.open(target, "w", encoding="utf-8") as fh:
    fh.write("ts_utc,tag_name,value,quality,quality_label,source,site,area,equipment" + NL)
    fh.write("old,row,1,192,GOOD,x,,," + NL)
fs.write_csv(target, {}, [{"gateway_id": "EM1", "ts_utc": "t1", "tag_name": "a",
                           "value": 1, "quality": 192, "quality_label": "GOOD",
                           "source": "power_modbus", "gateway_name": "EM1"}],
             lambda t: t)
rotated = [f for f in os.listdir(rt) if f != "export.csv"]
check("the mismatched file is moved aside", len(rotated) == 1, rotated)
with io.open(target, encoding="utf-8") as fh:
    head = fh.readline().strip().split(",")
check("  the new file carries the header we write",
      head == fs.CANONICAL_HEADER, head)
check("  and the old data is still on disk", bool(rotated) and
      os.path.getsize(os.path.join(rt, rotated[0])) > 0,
      "rotating must never mean deleting")

print()
print("[a wide header is REUSED, not rebuilt on every batch]")
# 2026-09-04, reported: "check why 2 csv files are being exported" - there
# were dozens, one roughly every second. A wide export with no ticked tags
# computes its columns from the tags PRESENT in the batch, and the insight.*
# tags are not produced every cycle. The header changed on nearly every write
# and the rotation fired each time.
wt = tempfile.mkdtemp(prefix="tn-wide-")
wide_path = os.path.join(wt, "wide.csv")


def wbatch(tags, ts):
    return [{"gateway_id": "EM1", "ts_utc": ts, "tag_name": t, "value": 1.0}
            for t in tags]


wsink = {"layout": "wide"}
fs.write_csv(wide_path, wsink, wbatch(["a", "b", "insight.x"], "t1"), lambda t: t)
fs.write_csv(wide_path, wsink, wbatch(["a", "b"], "t2"), lambda t: t)
fs.write_csv(wide_path, wsink, wbatch(["a", "b"], "t3"), lambda t: t)
check("a batch missing a column does NOT start a new file",
      len(os.listdir(wt)) == 1, os.listdir(wt))
with io.open(wide_path, encoding="utf-8") as fh:
    wlines = [ln.strip() for ln in fh if ln.strip()]
check("  the original header is kept", wlines[0].endswith("insight.x"), wlines[0])
check("  and the absent tag is blank in that row", wlines[2].endswith(","), wlines[2])
fs.write_csv(wide_path, wsink, wbatch(["a", "b", "brand_new"], "t4"), lambda t: t)
check("a genuinely NEW tag still rotates", len(os.listdir(wt)) == 2,
      "a column the existing file cannot hold is the one case that justifies "
      "a new file")

print()
print("[every file engine can start a new file, not just Excel]")
check("weekly and size are offered",
      "weekly" in fs.ROLLOVERS and "size" in fs.ROLLOVERS, fs.ROLLOVERS)
import datetime as _dt                                   # noqa: E402
_when = _dt.datetime(2026, 9, 4, 7, 30)
check("  weekly names the ISO week",
      fs.rollover_path("/x/e.csv", "weekly", _when).endswith("e_2026-W36.csv"),
      fs.rollover_path("/x/e.csv", "weekly", _when))
check("  size keeps ONE live path", fs.rollover_path("/x/e.csv", "size", _when) == "/x/e.csv",
      "a size limit cannot be expressed as a period in the name")
st = tempfile.mkdtemp(prefix="tn-size-")
sp = os.path.join(st, "big.csv")
with io.open(sp, "w", encoding="utf-8") as fh:
    fh.write("x" * 5000)
check("a file over its limit is moved aside",
      bool(fs.rotate_if_too_big(sp, 1000)) and len(os.listdir(st)) == 1,
      "renamed with a timestamp; nothing deleted")
check("  and one under it is left alone",
      fs.rotate_if_too_big(sp, 10 ** 9) == "")
csv_src = io.open(os.path.join(ROOT, "backend", "app", "services", "file_sinks.py"),
                  encoding="utf-8", errors="replace").read()
# Behaviour, not the exact call expression: this used to match a literal that
# changed the moment the path started being resolved, which made a passing
# feature look broken.
_rt = tempfile.mkdtemp(prefix="tn-roll-")
_row = [{"gateway_id": "EM1", "ts_utc": "2026-09-29 10:00:00", "tag_name": "a",
         "value": 1.0, "quality": 192, "quality_label": "GOOD", "source": "x"}]
for _engine, _writer, _name in (("csv", fs.write_csv, "r.csv"),
                                ("txt", fs.write_txt, "r.txt")):
    _p = os.path.join(_rt, _name)
    _writer(_p, {"rollover": "daily"}, _row, lambda t: t)
    _made = [f for f in os.listdir(_rt) if f.startswith(_name.split(".")[0])]
    check("%s honours rollover, not just xlsx" % _engine.upper(),
          any("_20" in f for f in _made), _made)

print()
print("[whatever the operator typed becomes a real, writable path]")
# 2026-09-29, reported: setting a new export location on another machine
# stopped the logging. The meter path used the string verbatim while the PLC
# path resolved it, so the same value behaved differently on the two paths.
pt = tempfile.mkdtemp(prefix="tn-rp-")
_prev = os.environ.get("TRUSTNODE_DATA_DIR")
os.environ["TRUSTNODE_DATA_DIR"] = pt
try:
    check("a RELATIVE path lands in the data directory",
          fs.resolve_path("exports/mill.csv").startswith(os.path.abspath(pt)),
          "it used to land wherever the service happened to be running")
    check("a folder with a trailing separator gets a file name",
          fs.resolve_path(pt + os.sep, "auto.csv").endswith("auto.csv"),
          "the folder picker hands back a directory, which cannot be opened")
    check("  and so does an existing folder",
          fs.resolve_path(pt, "auto.csv").endswith("auto.csv"))
    check("quotes from Explorer's Copy as path are stripped",
          not fs.resolve_path('"%s"' % os.path.join(pt, "q.csv")).count('"'),
          "a quoted path was taken literally and never opened")
    check("the parent directory is created",
          os.path.isdir(os.path.dirname(fs.resolve_path("deep/er/x.csv"))),
          "a new location must work the moment it is saved")
finally:
    if _prev is None:
        os.environ.pop("TRUSTNODE_DATA_DIR", None)
    else:
        os.environ["TRUSTNODE_DATA_DIR"] = _prev

print()
print("[a TXT export is readable - it has a header too]")
tpath = os.path.join(rt, "export.txt")
fs.write_txt(tpath, {}, [{"gateway_id": "EM1", "ts_utc": "t1", "tag_name": "a",
                          "value": 1, "quality_label": "GOOD"}], lambda t: t)
with io.open(tpath, encoding="utf-8") as fh:
    tlines = [ln.strip() for ln in fh if ln.strip()]
check("the TXT names its columns", tlines[0].split("|") == fs.CANONICAL_HEADER,
      "nine unlabelled fields is not an export anybody can read")
check("  and then the reading", tlines[1].split("|")[4] == "a", tlines[1])

print()
print("[an export with no gateway scope covers EVERY source]")
check("empty gateway_filters attaches the sink",
      "!scope.length || scope.map(String).includes(deviceId)" in app,
      "an unscoped export - the DEFAULT - attached to nothing, so the TXT "
      "file stayed empty while its configuration looked correct")
check("the meter status is polled beyond the Power pages",
      "METER_STATE_PAGES" in app and "wantsFastPower" in app,
      "the gateway row and the footer showed whatever was last fetched")

print()
print("[the two hand-written allowlists this has to survive]")
# The trap this codebase has paid for repeatedly: a field the dialog collects
# is dropped on save, or dropped on its way to the worker, and the feature
# looks broken while the config looks right.
saved = app.split("const finalizeSave")[1][:4000] if "const finalizeSave" in app else ""
to_sink = app.split("const toSink = (conn)")[1][:2500] if "const toSink = (conn)" in app else ""
for field in ("export_interval_s", "rollover", "layout", "max_file_mb"):
    check("%-18s survives the SAVE list" % field, field in saved,
          "collected by the dialog and then silently discarded")
    check("%-18s reaches the WORKER payload" % field, field in to_sink,
          "saved correctly and then never sent - identical symptom, different "
          "allowlist")

print()
print("RESULT: {0}".format("PASS" if not FAILS else "FAIL - " + ", ".join(FAILS)))
sys.exit(0 if not FAILS else 2)
