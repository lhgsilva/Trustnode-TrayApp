# -*- coding: utf-8 -*-
"""A running gateway actually produces CSV, TXT and Excel files.

2026-09-03, asked directly: "have you tested if gateway trigger condition is
made and if excel is being logged?"

The honest answer at the time was no. `test_file_exports.py` exercises the
writer functions in isolation - filters, rate, rollover, the atomic save - and
they all passed while the operator's export file had never received a single
row. Unit-testing a writer proves the writer works; it does not prove anything
reaches it. That is exactly the gap that let a broken export ship.

So this test runs the whole path: a Modbus simulator, a real backend, a real
gateway with csv_file, txt_file and xlsx_file sinks attached, and then it
opens the files off disk and counts what is in them.

It also asserts the two halves TOGETHER, which is the part that matters to an
operator: while a collection trigger is unsatisfied the files must not grow,
and when the condition is met they must start growing again. An export that
keeps writing through a pause would put data in a file that the historian
deliberately refused to keep.
"""
from __future__ import annotations

import json
import os
import socket
import subprocess
import sys
import tempfile
import threading
import time
import urllib.error
import urllib.request

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "backend"))
os.environ.setdefault("TRUSTNODE_SKIP_DOTENV", "1")

FAILS: list[str] = []


def check(name, ok, detail=""):
    print("  {0:56s}: {1}{2}".format(name, "PASS" if ok else "FAIL",
                                     (" - " + str(detail)[:150]) if detail else ""))
    if not ok:
        FAILS.append(name)


def free_port():
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


PORT = str(free_port())
API = "http://127.0.0.1:" + PORT
MB_PORT = free_port()
GW = "gw-export-e2e"
TAG = "Level"

from pymodbus.datastore import (ModbusSequentialDataBlock,       # noqa: E402
                                ModbusServerContext, ModbusSlaveContext)
from pymodbus.server import StartTcpServer                        # noqa: E402

_block = ModbusSequentialDataBlock(0, [0] * 200)
_ctx = ModbusServerContext(slaves=ModbusSlaveContext(
    hr=_block, ir=_block,
    di=ModbusSequentialDataBlock(0, [0] * 100),
    co=ModbusSequentialDataBlock(0, [0] * 100), zero_mode=True), single=True)
threading.Thread(
    target=lambda: StartTcpServer(context=_ctx, address=("127.0.0.1", MB_PORT)),
    daemon=True).start()
time.sleep(1.5)


def set_register(value: int) -> None:
    _block.setValues(0, [int(value)])


tmp = tempfile.mkdtemp(prefix="tn-exp-")
out_dir = os.path.join(tmp, "exports")
os.makedirs(out_dir, exist_ok=True)
CSV_PATH = os.path.join(out_dir, "export.csv")
TXT_PATH = os.path.join(out_dir, "export.txt")
XLSX_PATH = os.path.join(out_dir, "export.xlsx")

env = dict(os.environ, TRUSTNODE_SKIP_DOTENV="1", TRUSTNODE_DATA_DIR=tmp,
           TRUSTNODE_APP_STORE_PATH=os.path.join(tmp, "s.db"),
           TRUSTNODE_BOOT_INTEGRITY_CHECK="never", TRUSTNODE_PORT=PORT)
log = open(os.path.join(tmp, "o.log"), "w")
_exe = os.environ.get("TRUSTNODE_TEST_EXE", "").strip()
if _exe:
    proc = subprocess.Popen([_exe], env=env, stdout=log, stderr=subprocess.STDOUT)
else:
    proc = subprocess.Popen([sys.executable, "-m", "app"],
                            cwd=os.path.join(ROOT, "backend"), env=env,
                            stdout=log, stderr=subprocess.STDOUT)

up = False
for _ in range(80):
    try:
        urllib.request.urlopen(API + "/api/health", timeout=3).read()
        up = True
        break
    except Exception:
        time.sleep(2)


def call(method, path, token=None, body=None, timeout=60):
    headers = {"Content-Type": "application/json"}
    if token:
        headers["Authorization"] = "Bearer " + token
    data = json.dumps(body).encode() if body is not None else None
    req = urllib.request.Request(API + path, data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return r.status, json.loads(r.read().decode() or "null")
    except urllib.error.HTTPError as e:
        return e.code, None
    except Exception as e:
        return 0, str(e)[:140]


def finish(code):
    try:
        proc.terminate()
        proc.wait(timeout=20)
    except Exception:
        proc.kill()
    sys.exit(code)


def line_rows(path: str) -> int:
    """Non-empty lines in a text export."""
    if not os.path.exists(path):
        return 0
    with open(path, encoding="utf-8", errors="replace") as fh:
        return sum(1 for line in fh if line.strip())


def csv_rows() -> int:
    return line_rows(CSV_PATH)


def txt_rows() -> int:
    return line_rows(TXT_PATH)


def xlsx_path_today() -> str:
    from app.services import file_sinks as fs
    return fs.rollover_path(XLSX_PATH, "daily")


def xlsx_rows() -> int:
    """Data rows in today's workbook, excluding the header."""
    path = xlsx_path_today()
    if not os.path.exists(path) or os.path.getsize(path) == 0:
        return 0
    from openpyxl import load_workbook
    # CLOSE the workbook. read_only mode keeps the file handle open on
    # Windows, and this is polled once a second - the leaked handles made
    # os.replace() in the writer fail, so the file stopped growing and the
    # TEST looked like a product bug. Reading a file must not change it.
    book = None
    try:
        book = load_workbook(path, read_only=True)
        return max(0, book.active.max_row - 1)
    except Exception:
        return 0
    finally:
        try:
            if book is not None:
                book.close()
        except Exception:
            pass


print("TrustNode - a running gateway writes CSV, TXT and Excel")
print()
check("the app started", up)
if not up:
    finish(2)

tok = (call("POST", "/api/auth/login",
            body={"username": "admin", "password": "admin"})[1] or {}).get("token")
check("admin login", bool(tok))
if not tok:
    finish(2)

CSV_SINK = {
    "id": "sink-csv", "name": "CSV Export", "engine": "csv_file",
    "file_path": CSV_PATH, "tag_filters": [], "gateway_filters": [],
    "csv_format": "", "csv_header": "", "export_interval_s": 0,
    "rollover": "none", "table": "plc_readings", "schema": "public",
}
TXT_SINK = {
    "id": "sink-txt", "name": "TXT Export", "engine": "txt_file",
    "file_path": TXT_PATH, "tag_filters": [], "gateway_filters": [],
    "csv_format": "", "csv_header": "", "export_interval_s": 0,
    "rollover": "none", "table": "plc_readings", "schema": "public",
}
XLSX_SINK = {
    "id": "sink-xlsx", "name": "Excel Export", "engine": "xlsx_file",
    "file_path": XLSX_PATH, "tag_filters": [], "gateway_filters": [],
    "export_interval_s": 0, "rollover": "daily",
    "table": "plc_readings", "schema": "public",
}
CONFIG = {
    "gateway_type": "modbus_tcp", "name": "Export E2E",
    "plc_ip": "127.0.0.1", "modbus_port": MB_PORT, "modbus_unit_id": 1,
    "modbus_registers": [{"name": TAG, "address": "4x:0", "function": "holding",
                          "kind": "uint16", "enabled": True}],
    "tags": [TAG], "interval_ms": 500,
    "equipment": "SIM", "site": "B", "area": "T",
}

set_register(75)
st, body = call("POST", "/api/plc/gateways/start", tok, {
    "gateway_id": GW, "config": CONFIG,
    "db_sink": CSV_SINK, "db_sinks": [CSV_SINK, TXT_SINK, XLSX_SINK]})
check("the gateway starts with CSV, TXT and Excel sinks", st == 200, body)

print()
print("[the files are actually produced on disk]")
deadline = time.time() + 45
while time.time() < deadline and (csv_rows() < 3 or txt_rows() < 3
                                 or xlsx_rows() < 3):
    time.sleep(1)

csv_n, txt_n, xlsx_n = csv_rows(), txt_rows(), xlsx_rows()
check("the CSV file exists", os.path.exists(CSV_PATH), CSV_PATH)
check("  and has data rows", csv_n >= 3, "%d row(s)" % csv_n)
check("the TXT file exists", os.path.exists(TXT_PATH), TXT_PATH)
check("  and has data rows", txt_n >= 3, "%d row(s)" % txt_n)
check("the EXCEL workbook exists", os.path.exists(xlsx_path_today()),
      xlsx_path_today())
check("  and has data rows", xlsx_n >= 3,
      "%d row(s) - the writer passed its unit test while the operator's file "
      "stayed empty, so only reading the real file proves anything" % xlsx_n)

if xlsx_n:
    from openpyxl import load_workbook
    _book = load_workbook(xlsx_path_today(), read_only=True)
    sheet = _book.active
    header = [c.value for c in next(sheet.iter_rows(min_row=1, max_row=1))]
    row2 = [c.value for c in next(sheet.iter_rows(min_row=2, max_row=2))]
    _book.close()
    from app.services import file_sinks as fs
    check("  the workbook carries the canonical header",
          header == fs.CANONICAL_HEADER, header)
    check("  and a real reading, not an empty row",
          any(v not in (None, "") for v in row2), row2)
    check("  naming the tag that was collected",
          TAG in [str(v) for v in row2], row2)

print()
print("[an unsatisfied trigger stops BOTH files]")
set_register(10)                       # below the threshold
st, _ = call("PUT", "/api/app-store/domain", tok, {
    "domain": "triggers_limits", "actor": "test",
    "payload": {"collection_triggers": [
        {"gateway_id": GW, "tag_name": TAG, "operator": ">=", "value": 50,
         "trigger_type": "continuous", "enabled": True, "kind": "tag"}],
        "collection_trigger_mode": "any", "trigger_rules": []}})
check("the trigger was saved", st == 200, st)
time.sleep(6)
csv_paused, txt_paused, xlsx_paused = csv_rows(), txt_rows(), xlsx_rows()
time.sleep(8)
check("the CSV stops growing while paused",
      csv_rows() == csv_paused, "%d row(s) leaked" % (csv_rows() - csv_paused))
check("the TXT stops growing while paused",
      txt_rows() == txt_paused, "%d row(s) leaked" % (txt_rows() - txt_paused))
check("the EXCEL workbook stops growing while paused",
      xlsx_rows() == xlsx_paused, "%d row(s) leaked" % (xlsx_rows() - xlsx_paused))

print()
print("[and the condition being MET starts them again]")
set_register(75)                       # condition true
deadline = time.time() + 40
while time.time() < deadline and (csv_rows() <= csv_paused
                                 or txt_rows() <= txt_paused
                                 or xlsx_rows() <= xlsx_paused):
    time.sleep(1)
check("the CSV resumes", csv_rows() > csv_paused,
      "%d -> %d" % (csv_paused, csv_rows()))
check("the TXT resumes", txt_rows() > txt_paused,
      "%d -> %d" % (txt_paused, txt_rows()))
check("the EXCEL workbook resumes", xlsx_rows() > xlsx_paused,
      "%d -> %d" % (xlsx_paused, xlsx_rows()))

print()
print("[a POWER METER exports too - the operator's actual case]")
# 2026-09-03, reported: "export to csv, txt and excel not working when gateway
# is collecting". The gateway path above was already proven; the METER path
# was not, and it was broken in two more places:
#
#   * _normalize_device builds a device field by field, so the `sinks` the UI
#     pushes were silently discarded on the way into the running meter;
#   * _load_config read `config_documents`, the legacy unscoped table, so the
#     meter ran on a copy of its configuration from the previous evening and
#     never saw the sinks at all.
#
# Neither shows up in a unit test of the writer, and neither shows up in the
# gateway test above. Only driving a real meter does.
import struct                                        # noqa: E402

MCSV = os.path.join(out_dir, "meter.csv")
MTXT = os.path.join(out_dir, "meter.txt")
MXLSX = os.path.join(out_dir, "meter.xlsx")


def meter_sink(engine, path, rollover="none"):
    return {"id": "m-" + engine, "name": "Meter " + engine, "engine": engine,
            "file_path": path, "tag_filters": [], "gateway_filters": ["EM-TEST"],
            "csv_format": "", "csv_header": "", "export_interval_s": 0,
            "rollover": rollover, "table": "plc_readings", "schema": "public"}


# 230.0 V as float32 big-endian in input registers 0-1, where 30001 resolves.
_hi, _lo = struct.unpack(">HH", struct.pack(">f", 230.0))
_block.setValues(0, [_hi, _lo])

st, _ = call("PUT", "/api/app-store/domain", tok, {
    "domain": "power_management_config", "actor": "test",
    "payload": {
        "enabled": True, "selected_device_id": "EM-TEST",
        "energy_price_eur_kwh": 0.2,
        "devices": [{
            "id": "EM-TEST", "name": "EM-TEST", "enabled": True,
            "ip": "127.0.0.1", "port": MB_PORT, "unit_id": 1,
            "poll_interval_ms": 1000,
            "registers": {"voltage_v": 30001},
            "register_scales": {"voltage_v": 1},
            "database_id": "",
            "sinks": [meter_sink("csv_file", MCSV),
                      meter_sink("txt_file", MTXT),
                      meter_sink("xlsx_file", MXLSX, "daily")],
        }],
    }})
check("the meter configuration was saved", st == 200, st)


def mxlsx_today():
    from app.services import file_sinks as fs
    return fs.rollover_path(MXLSX, "daily")


def mxlsx_rows():
    path = mxlsx_today()
    if not os.path.exists(path) or os.path.getsize(path) == 0:
        return 0
    from openpyxl import load_workbook
    book = None
    try:
        book = load_workbook(path, read_only=True)
        return max(0, book.active.max_row - 1)
    except Exception:
        return 0
    finally:
        try:
            if book is not None:
                book.close()
        except Exception:
            pass


deadline = time.time() + 60
while time.time() < deadline and (line_rows(MCSV) < 2 or line_rows(MTXT) < 2
                                  or mxlsx_rows() < 2):
    time.sleep(2)

check("the meter writes a CSV", line_rows(MCSV) >= 2,
      "%d row(s) - the pushed sinks were dropped by _normalize_device, and the "
      "config came from a stale legacy table" % line_rows(MCSV))
check("the meter writes a TXT", line_rows(MTXT) >= 2, "%d row(s)" % line_rows(MTXT))
check("the meter writes an Excel workbook", mxlsx_rows() >= 2,
      "%d row(s) at %s" % (mxlsx_rows(), mxlsx_today()))

print()
print("RESULT: {0}".format("PASS" if not FAILS else "FAIL - " + ", ".join(FAILS)))
finish(0 if not FAILS else 2)
