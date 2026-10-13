# -*- coding: utf-8 -*-
"""CSV / TXT / XLSX export sinks: what to write, and how often.

2026-09-02, requested: "when setting a parallel database logging we should be
able to pick different gateways running, selecting them, and all the tags we
want to export, and also the rate. At the moment this is not working."

Three separate things were missing, and this module holds the two that belong
to the backend.

SCOPE - which rows a sink accepts

A sink carries `gateway_filters` (source ids) and `tag_filters` (tag names).
Empty means "everything", which is how every sink written before this existed
behaved, so nothing changes for them. The PLC path already honoured both; the
power-meter path did not, so a meter wrote all forty of its registers to a CSV
no matter what the operator ticked.

RATE - how often it writes

Without a rate a sink writes on every poll cycle. A 1 s meter with 40
registers is 40 rows a second, 3.4 million rows a day, which is not an export
anybody reads. `export_interval_s` keeps the LATEST value per tag and releases
one row per tag per interval. Latest rather than first: an export sampled at
60 s should show what the tag reads now, not what it read 59 s ago.

A rate is a SAMPLE, not an aggregate. It does not average, and it does not
claim to - a mean would hide exactly the spike an operator opens the file to
find.

XLSX - why it rolls over

.xlsx is a zip archive, so unlike CSV it cannot be appended to; adding a row
means rewriting the workbook. Left alone that gets slower as the file grows
and can corrupt the whole history if the app stops mid-write. So an xlsx sink
rolls over - one workbook per hour, day or month - and the export rate bounds
how often each one is rewritten. Together they keep the rewrite small and
losable to at most one period.
"""
from __future__ import annotations

import os
import time as _time
from datetime import datetime
from typing import Any, Dict, Iterable, List

CANONICAL_HEADER = ["ts_local", "ts_utc", "gateway_id", "gateway_name",
                    "tag_name", "value", "quality", "quality_label", "source"]

ROLLOVERS = ("none", "hourly", "daily", "weekly", "monthly", "size")

# sink_key -> {"last": monotonic, "pending": {(gateway_id, tag_name): row}}
_rate_state: Dict[str, Dict[str, Any]] = {}


def filter_rows(sink: Dict[str, Any], rows: Iterable[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Drop rows this sink was not asked to export.

    An empty filter list means "accept everything" - configs written before
    scoping existed must keep working.
    """
    gateways = {str(g or "").strip() for g in (sink.get("gateway_filters") or []) if g}
    tags = {str(t or "").strip() for t in (sink.get("tag_filters") or []) if t}
    if not gateways and not tags:
        return list(rows)
    kept: List[Dict[str, Any]] = []
    for r in rows:
        if gateways and str(r.get("gateway_id") or "").strip() not in gateways:
            continue
        if tags and str(r.get("tag_name") or "").strip() not in tags:
            continue
        kept.append(r)
    return kept


def export_interval_s(sink: Dict[str, Any]) -> float:
    """The sink's export rate in seconds. 0 (the default) = every cycle."""
    try:
        return max(0.0, float(sink.get("export_interval_s") or 0))
    except (TypeError, ValueError):
        return 0.0


def due_rows(sink_key: str, sink: Dict[str, Any],
             rows: Iterable[Any], now: float | None = None,
             key_of=None) -> List[Any]:
    """Return the rows this sink should write RIGHT NOW.

    With no rate configured every row passes straight through, so an existing
    sink behaves exactly as before. With a rate, the latest row per
    (gateway, tag) is held and released once per interval.

    `key_of` lets the PLC path pass its GatewayReading objects through the
    same throttle as the meter path's dicts - one implementation, so the two
    cannot drift into behaving differently.
    """
    interval = export_interval_s(sink)
    rows = list(rows)
    if interval <= 0:
        return rows
    keyed = key_of or (lambda r: (str(r.get("gateway_id") or ""),
                                  str(r.get("tag_name") or "")))
    moment = _time.monotonic() if now is None else now
    state = _rate_state.setdefault(str(sink_key), {"last": 0.0, "pending": {}})
    for r in rows:
        state["pending"][keyed(r)] = r
    if moment - float(state["last"] or 0.0) < interval:
        return []
    state["last"] = moment
    released = list(state["pending"].values())
    state["pending"] = {}
    return released


def reset_rate_state(sink_key: str | None = None) -> None:
    """Forget throttling state - used by tests and when a sink is deleted."""
    if sink_key is None:
        _rate_state.clear()
    else:
        _rate_state.pop(str(sink_key), None)


def rollover_path(file_path: str, rollover: str, now: datetime | None = None) -> str:
    """Where this period's file lives.

    "none" keeps the configured path. Anything else inserts the period before
    the extension, so `power.xlsx` becomes `power_2026-09-02.xlsx` - sorting
    the directory by name puts the periods in order.
    """
    mode = str(rollover or "none").strip().lower()
    if mode not in ROLLOVERS or mode == "none":
        return file_path
    if mode == "size":
        # Size rollover keeps ONE live path and renames it aside when it grows
        # past the limit - see rotate_if_too_big(). A period cannot name it.
        return file_path
    moment = now or datetime.now()
    stamp = {"hourly": "%Y-%m-%d_%H", "daily": "%Y-%m-%d",
             # ISO week, so a file never straddles a year boundary ambiguously.
             "weekly": "%G-W%V", "monthly": "%Y-%m"}[mode]
    base, ext = os.path.splitext(file_path)
    return "%s_%s%s" % (base, moment.strftime(stamp), ext)


def max_file_bytes(sink: Dict[str, Any]) -> int:
    """Size limit in bytes for a `size` rollover. 0 = no limit."""
    try:
        mb = float(sink.get("max_file_mb") or 0)
    except (TypeError, ValueError):
        return 0
    return int(max(0.0, mb) * 1024 * 1024)


def rotate_if_too_big(path: str, limit_bytes: int) -> str:
    """Move the file aside once it passes its size limit.

    Requested 2026-09-04: "option to create a new file every hour, day, week
    or file size". A size limit is the one an operator reaches for when the
    rate is unpredictable - a busy day should not produce a file too big to
    open just because it happened to fall inside one period.
    """
    if limit_bytes <= 0 or not path:
        return ""
    try:
        if not os.path.exists(path) or os.path.getsize(path) < limit_bytes:
            return ""
        stamp = datetime.now().strftime("%Y-%m-%dT%H-%M-%S")
        base, ext = os.path.splitext(path)
        moved = "%s.%s%s" % (base, stamp, ext)
        os.replace(path, moved)
        return moved
    except Exception:
        return ""


def row_values(row: Dict[str, Any], ts_local: str) -> List[Any]:
    """One row in canonical column order."""
    return [
        ts_local,
        str(row.get("ts_utc") or ""),
        str(row.get("gateway_id") or ""),
        str(row.get("gateway_name") or ""),
        str(row.get("tag_name") or ""),
        row.get("value"),
        row.get("quality"),
        str(row.get("quality_label") or ""),
        str(row.get("source") or ""),
    ]


def write_xlsx(file_path: str, rows: List[Dict[str, Any]], rollover: str,
               ts_local_of, now: datetime | None = None,
               max_file_mb: float = 0) -> str:
    """Append rows to this period's workbook, creating it if needed.

    Returns the path actually written, so a caller can log which period the
    rows landed in.
    """
    if not rows:
        return ""
    from openpyxl import Workbook, load_workbook          # local: heavy import

    path = rollover_path(resolve_path(file_path, "trustnode_export.xlsx"),
                         rollover, now)
    rotate_if_too_big(path, max_file_bytes({"max_file_mb": max_file_mb}))
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    if os.path.exists(path) and os.path.getsize(path) > 0:
        book = load_workbook(path)
        sheet = book.active
    else:
        book = Workbook()
        sheet = book.active
        sheet.title = "readings"
        sheet.append(CANONICAL_HEADER)
    for r in rows:
        sheet.append(row_values(r, ts_local_of(str(r.get("ts_utc") or ""))))
    # Write to a temporary file and replace, so a crash mid-save cannot leave
    # a half-written zip where the previous good workbook was.
    tmp = path + ".part"
    book.save(tmp)
    os.replace(tmp, path)
    return path


# ---------------------------------------------------------------- layout
#
# 2026-09-03, requested: "add an option to the data to be exported when the
# first column is the timestamp and the columns different tags, so we can get
# multiple values per timestamp".
#
#   long (default)  one row per reading - what every existing sink writes.
#   wide            one row per timestamp, one COLUMN per tag. This is the
#                   shape a spreadsheet wants: forty tags become forty
#                   columns instead of forty rows a second.
#
# The header DEFINES the columns of a wide file, so the tag set has to stay
# fixed for its life. When a tag appears that the header does not name, the
# file is rotated rather than written into: silently dropping the new tag
# loses data, and appending it under the wrong column is worse.

LAYOUTS = ("long", "wide")


def layout_of(sink: Dict[str, Any]) -> str:
    value = str(sink.get("layout") or "long").strip().lower()
    return value if value in LAYOUTS else "long"


def tag_columns(sink: Dict[str, Any], rows: Iterable[Dict[str, Any]]) -> List[str]:
    """Column order for a wide file.

    An explicit tag_filters list wins, because it is the operator's own
    ordering and it is stable across restarts. Otherwise the tags present,
    sorted, so the same data always produces the same file.
    """
    chosen = [str(t or "").strip() for t in (sink.get("tag_filters") or []) if str(t or "").strip()]
    if chosen:
        return chosen
    return sorted({str(r.get("tag_name") or "").strip() for r in rows
                   if str(r.get("tag_name") or "").strip()})


def header_for(sink: Dict[str, Any], rows: Iterable[Dict[str, Any]]) -> List[str]:
    if layout_of(sink) == "wide":
        return ["ts_local", "ts_utc", "gateway_id"] + tag_columns(sink, rows)
    return list(CANONICAL_HEADER)


def shape_rows(sink: Dict[str, Any], rows: List[Dict[str, Any]],
               ts_local_of, header: List[str] | None = None) -> List[List[Any]]:
    """Rows as they will appear in the file, in the sink's layout."""
    if layout_of(sink) != "wide":
        return [row_values(r, ts_local_of(str(r.get("ts_utc") or ""))) for r in rows]
    cols = (header or header_for(sink, rows))[3:]
    grouped: Dict[tuple, Dict[str, Any]] = {}
    order: List[tuple] = []
    for r in rows:
        key = (str(r.get("gateway_id") or ""), str(r.get("ts_utc") or ""))
        if key not in grouped:
            grouped[key] = {}
            order.append(key)
        grouped[key][str(r.get("tag_name") or "")] = r.get("value")
    out: List[List[Any]] = []
    for key in order:
        gw, ts = key
        values = grouped[key]
        # A tag with no reading in this cycle stays EMPTY rather than 0 - a
        # blank cell and a real zero must not look the same.
        out.append([ts_local_of(ts), ts, gw] + [values.get(c, "") for c in cols])
    return out


def existing_header(path: str, delimiter: str = ",") -> List[str] | None:
    """The header a text export already carries, or None if it has none."""
    try:
        if not os.path.exists(path) or os.path.getsize(path) == 0:
            return None
        with open(path, encoding="utf-8", errors="replace") as handle:
            first = handle.readline().strip()
    except Exception:
        return None
    if not first:
        return None
    return [c.strip() for c in first.split(delimiter)]


def reconcile_header(path: str, header: List[str],
                     delimiter: str = ",") -> List[str]:
    """The header to write, preferring the one the file already has.

    2026-09-04: a wide export with no ticked tags computes its columns from
    the tags PRESENT in the batch, and batches vary - the insight.* tags are
    not produced every cycle. So the header changed on nearly every write, the
    rotation below fired each time, and the operator got a new CSV roughly
    every second instead of one file.

    A file's header IS its column set. If everything in this batch fits the
    columns already there, keep them and let the missing ones render empty -
    that is exactly what a wide row does for a tag with no reading. Only a
    genuinely NEW tag, which cannot be represented at all, justifies starting
    a new file.
    """
    current = existing_header(path, delimiter)
    if not current or len(current) < 3 or current[:3] != [str(c) for c in header[:3]]:
        return header
    have = set(current[3:])
    if set(header[3:]) - have:
        return header          # a column the existing file cannot hold
    return current


def rotate_if_header_differs(path: str, header: List[str],
                             delimiter: str = ",") -> str:
    """Move a file aside when its header no longer matches what we write.

    2026-09-03: the operator opened a CSV whose header said `ts_utc,tag_name,
    value,...` while every row underneath was `ts_local,ts_utc,gateway_id,...`
    - the file had been created by one writer and appended to by another, and
    the columns had been meaningless ever since. Appending under a header we
    did not write is never right.

    The old file is renamed with a timestamp, so nothing is lost and neither
    file ends up with two column orders in it. Returns the rotated path, or ""
    when nothing needed rotating.
    """
    current = existing_header(path, delimiter)
    if current is None or current == [str(c) for c in header]:
        return ""
    stamp = datetime.now().strftime("%Y-%m-%dT%H-%M-%S")
    base, ext = os.path.splitext(path)
    moved = "%s.%s%s" % (base, stamp, ext)
    try:
        os.replace(path, moved)
        return moved
    except Exception:
        return ""


def write_csv(file_path: str, sink: Dict[str, Any], rows: List[Dict[str, Any]],
              ts_local_of) -> str:
    """Append rows to a CSV in the sink's layout, header first."""
    import csv as _csv
    if not rows or not file_path:
        return ""
    # 2026-09-04: rollover used to apply to Excel ALONE, so a CSV or TXT
    # export grew without limit until somebody noticed. Every file engine
    # honours it now - hourly/daily/weekly/monthly by period, or by size.
    path = rollover_path(resolve_path(file_path, "trustnode_export.csv"),
                         str(sink.get("rollover") or "none"))
    rotate_if_too_big(path, max_file_bytes(sink))
    header = reconcile_header(path, header_for(sink, rows), ",")
    rotate_if_header_differs(path, header, ",")
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    fresh = not os.path.exists(path) or os.path.getsize(path) == 0
    with open(path, "a", encoding="utf-8", newline="") as handle:
        writer = _csv.writer(handle)
        if fresh:
            writer.writerow(header)
        for values in shape_rows(sink, rows, ts_local_of, header):
            writer.writerow(values)
    return path


def write_txt(file_path: str, sink: Dict[str, Any], rows: List[Dict[str, Any]],
              ts_local_of, delimiter: str = "|") -> str:
    """Same content as the CSV, delimiter-separated, header first.

    A TXT export used to write bare rows with no header at all, so a file of
    nine unlabelled fields was left for the reader to decode. It carries the
    same header as the CSV now - the point of an export is that somebody else
    can read it.
    """
    if not rows or not file_path:
        return ""
    path = rollover_path(resolve_path(file_path, "trustnode_export.txt"),
                         str(sink.get("rollover") or "none"))
    rotate_if_too_big(path, max_file_bytes(sink))
    header = reconcile_header(path, header_for(sink, rows), delimiter)
    rotate_if_header_differs(path, header, delimiter)
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    fresh = not os.path.exists(path) or os.path.getsize(path) == 0
    with open(path, "a", encoding="utf-8") as handle:
        if fresh:
            handle.write(delimiter.join(str(c) for c in header) + "\n")
        for values in shape_rows(sink, rows, ts_local_of, header):
            handle.write(delimiter.join("" if v is None else str(v)
                                        for v in values) + "\n")
    return file_path


def resolve_path(file_path: str, fallback_name: str = "trustnode_export.csv") -> str:
    """Turn whatever the operator typed into a real, writable file path.

    2026-09-29, reported: setting a new export location on another machine
    stopped the logging. The meter path used the string verbatim while the PLC
    path resolved it, so the two behaved differently for the same value:

      * a RELATIVE path landed wherever the service happened to be running,
        which on a packaged install is not the data directory;
      * a path that names a FOLDER (from the folder picker, or typed with a
        trailing slash) could never be opened as a file;
      * a path pasted with quotes around it, as Explorer's "Copy as path"
        gives you, was taken literally.

    The parent directory is created here, so a new location works the moment
    it is saved rather than failing silently on the first write.
    """
    raw = str(file_path or "").strip().strip('"').strip("'")
    if not raw:
        raw = fallback_name
    if raw.endswith(("/", "\\")) or os.path.isdir(raw):
        raw = os.path.join(raw, fallback_name)
    if not os.path.isabs(raw):
        base = os.environ.get("TRUSTNODE_DATA_DIR", "").strip() or os.path.join(
            os.path.expanduser("~"), ".trustnode_edge", "data")
        raw = os.path.join(base, raw)
    full = os.path.abspath(raw)
    os.makedirs(os.path.dirname(full) or ".", exist_ok=True)
    return full
